"""SEC-audit hotfix (2026-02) : coherence Bilan HTML vs Bilan Excel.

Bug client LEFRANCQ : L'export Excel du bilan etait DESEQUILIBRE (Actif
25 891.78 != Passif 33 135.87) alors que la vue HTML montrait un bilan
equilibre a 16 844.09. Cause : l'export Excel utilisait une agregation
naive des `journal_entries` sans passer par `compute_bilan_data()`,
ignorant :
- Repartition du compte 499 (boni/mali)
- Exclusion des reversals
- Agregation par proprietaire (V.A / VI.A)
- View_mode before/after distribution

Test : deux totaux (Excel via endpoint) doivent egaler ceux du bilan HTML.
"""
import asyncio
import io
import os
import sys
import uuid

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import httpx
from openpyxl import load_workbook
from motor.motor_asyncio import AsyncIOMotorClient
from dotenv import load_dotenv

load_dotenv("/app/backend/.env")

API = "http://localhost:8001"
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "admin@copro.be")
ADMIN_PWD = os.environ["ADMIN_PASSWORD"]


async def _seed(db):
    tag = uuid.uuid4().hex[:8]
    copro_id = f"test-bilanexcel-{tag}"
    owner_id = f"owner-{tag}"

    await db.coproprietes.insert_one({
        "id": copro_id, "name": "TEST bilan excel",
        "bank_accounts": [{"iban": "BE12345678901234", "pcmn_number": "55000000",
                           "account_type": "vue", "label": "Bank"}],
    })
    await db.owners.insert_one({
        "id": owner_id, "name": "TEST Owner",
        "copropriete_id": copro_id,
        "tier_accounts": {copro_id: {"provisions": "40000010", "reserve": "40010010"}},
    })
    await db.lots.insert_one({
        "id": f"lot-{tag}", "copropriete_id": copro_id, "owner_id": owner_id,
        "number": "1", "quotity": 1000, "type": "appartement",
    })
    await db.pcmn_accounts.insert_many([
        {"copropriete_id": copro_id, "number": "40000010", "name": "Prov TEST", "active": True},
        {"copropriete_id": copro_id, "number": "40010010", "name": "Res TEST", "active": True},
        {"copropriete_id": copro_id, "number": "55000000", "name": "Banque", "active": True},
        {"copropriete_id": copro_id, "number": "160", "name": "Fonds de reserve general", "active": True},
    ])

    # AN d'ouverture 2025-01-01 : Owner debit 500 (provisions) + Fonds reserve credit 500
    await db.journal_entries.insert_one({
        "id": f"an-{tag}", "copropriete_id": copro_id,
        "date": "2025-01-01", "journal_type": "AN",
        "reference": "AN-2025-OPEN", "is_opening_balance": True,
        "description": "Report solde 2025",
        "lines": [
            {"account_number": "40000010", "third_party_id": owner_id,
             "debit": 500.0, "credit": 0.0},
            {"account_number": "160", "debit": 0.0, "credit": 500.0,
             "account_name": "Fonds de reserve general"},
        ],
    })
    # VE : appel 1000
    await db.journal_entries.insert_one({
        "id": f"ve-{tag}", "copropriete_id": copro_id,
        "date": "2025-07-01", "journal_type": "VE",
        "reference": "VE-2025-01",
        "description": "Appel",
        "lines": [
            {"account_number": "40000010", "third_party_id": owner_id,
             "debit": 1000.0, "credit": 0.0},
            {"account_number": "70000000", "debit": 0.0, "credit": 1000.0,
             "account_name": "Provisions appelees"},
        ],
    })
    return copro_id, owner_id


async def _cleanup(db, copro_id):
    await db.coproprietes.delete_one({"id": copro_id})
    await db.owners.delete_many({"copropriete_id": copro_id})
    await db.lots.delete_many({"copropriete_id": copro_id})
    await db.pcmn_accounts.delete_many({"copropriete_id": copro_id})
    await db.journal_entries.delete_many({"copropriete_id": copro_id})


async def main():
    client = AsyncIOMotorClient(os.environ["MONGO_URL"])
    db = client[os.environ["DB_NAME"]]

    copro_id, owner_id = await _seed(db)
    try:
        async with httpx.AsyncClient(base_url=API, timeout=30.0) as h:
            r = await h.post("/api/auth/login",
                             json={"email": ADMIN_EMAIL, "password": ADMIN_PWD})
            token = r.cookies.get("access_token")
            hdr = {"Authorization": f"Bearer {token}", "X-Copropriete-Id": copro_id}

            params = {"copropriete_id": copro_id, "date_to": "2025-12-31",
                      "view_mode": "before_distribution"}

            # HTML bilan
            r = await h.get("/api/reports/bilan", params=params, headers=hdr)
            assert r.status_code == 200, r.text
            html = r.json()
            html_actif = html["total_actif"]
            html_passif = html["total_passif"]
            print(f"HTML bilan : ACTIF={html_actif}, PASSIF={html_passif}, equilibre={html['equilibre']}")
            assert html["equilibre"], f"Bilan HTML desequilibre ecart={html['ecart']}"

            # Excel bilan
            r = await h.get("/api/exports/bilan.xlsx", params=params, headers=hdr)
            assert r.status_code == 200, r.text
            wb = load_workbook(io.BytesIO(r.content), data_only=True)
            ws = wb.active

            # Find TOTAL ACTIF and TOTAL PASSIF cells
            excel_actif = None
            excel_passif = None
            for row_idx in range(1, ws.max_row + 1):
                for col_idx in range(1, 4):
                    v = ws.cell(row=row_idx, column=col_idx).value
                    if v == "TOTAL ACTIF":
                        excel_actif = ws.cell(row=row_idx, column=3).value
                    elif v == "TOTAL PASSIF":
                        excel_passif = ws.cell(row=row_idx, column=3).value

            print(f"Excel bilan : ACTIF={excel_actif}, PASSIF={excel_passif}")
            assert excel_actif is not None, "TOTAL ACTIF not found in Excel"
            assert excel_passif is not None, "TOTAL PASSIF not found in Excel"
            assert abs(excel_actif - html_actif) < 0.01, f"Excel ACTIF {excel_actif} != HTML ACTIF {html_actif}"
            assert abs(excel_passif - html_passif) < 0.01, f"Excel PASSIF {excel_passif} != HTML PASSIF {html_passif}"
            assert abs(excel_actif - excel_passif) < 0.01, f"Excel DESEQUILIBRE: {excel_actif} != {excel_passif}"

            print(f"\nPASS : Excel = HTML = equilibre ({excel_actif} = {excel_passif})")

    finally:
        await _cleanup(db, copro_id)
        client.close()


if __name__ == "__main__":
    asyncio.run(main())
